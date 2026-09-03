#!/usr/bin/env python3
"""Build the shareable Russian GZWhisper product-functionality brief."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "GZWHISPER_PRODUCT_FUNCTIONALITY_RU.md"
OUTPUT = ROOT / "docs" / "GZWHISPER_PRODUCT_FUNCTIONALITY_RU.docx"

INK = RGBColor(0x0B, 0x25, 0x45)
BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
MUTED = RGBColor(0x5D, 0x68, 0x75)
LIGHT_FILL = "F4F6F9"


def set_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_or_paragraph_shading(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shading = p_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        p_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_left_border(paragraph, color="2E74B5", size="18", space="8"):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), size)
    left.set(qn("w:space"), space)
    left.set(qn("w:color"), color)
    borders.append(left)


def add_page_field(paragraph):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.extend([begin, instr, separate, text, end])
    set_font(run, size=9, color=MUTED)


def configure_styles(doc):
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0x18, 0x1D, 0x24)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.widow_control = True

    style_specs = {
        "Title": (30, INK, 0, 8),
        "Subtitle": (14, MUTED, 0, 14),
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in style_specs.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.widow_control = True

    if "Cover Title" not in styles:
        cover_title = styles.add_style("Cover Title", WD_STYLE_TYPE.PARAGRAPH)
    else:
        cover_title = styles["Cover Title"]
    cover_title.base_style = normal
    cover_title.font.name = "Calibri"
    cover_title._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    cover_title._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    cover_title._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    cover_title.font.size = Pt(30)
    cover_title.font.color.rgb = INK
    cover_title.font.bold = True
    cover_title.paragraph_format.space_before = Pt(0)
    cover_title.paragraph_format.space_after = Pt(8)
    cover_title.paragraph_format.keep_with_next = True
    cover_title.paragraph_format.widow_control = True

    if "Research Prompt" not in styles:
        quote = styles.add_style("Research Prompt", WD_STYLE_TYPE.PARAGRAPH)
    else:
        quote = styles["Research Prompt"]
    quote.base_style = normal
    quote.font.name = "Calibri"
    quote._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    quote._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    quote._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    quote.font.size = Pt(10.5)
    quote.paragraph_format.left_indent = Inches(0.22)
    quote.paragraph_format.right_indent = Inches(0.12)
    quote.paragraph_format.space_before = Pt(2)
    quote.paragraph_format.space_after = Pt(4)
    quote.paragraph_format.line_spacing = 1.18
    quote.paragraph_format.widow_control = True

    if "TOC Entry" not in styles:
        toc = styles.add_style("TOC Entry", WD_STYLE_TYPE.PARAGRAPH)
    else:
        toc = styles["TOC Entry"]
    toc.base_style = normal
    toc.font.name = "Calibri"
    toc._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    toc._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    toc._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    toc.font.size = Pt(10.5)
    toc.font.color.rgb = DARK_BLUE
    toc.paragraph_format.space_after = Pt(3)
    toc.paragraph_format.line_spacing = 1.12


def create_numbering(doc, kind):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    level.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    level.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    level.append(lvl_jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "260")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    level.append(p_pr)
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    r_pr.append(fonts)
    level.append(r_pr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])
    p_pr.append(num_pr)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25


def add_inline_markdown(paragraph, text, default_size=None, default_color=None):
    pattern = re.compile(r"(\*\*.+?\*\*|`.+?`)")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_font(run, size=default_size, color=default_color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_font(run, size=default_size, color=default_color, bold=True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_font(run, name="Menlo", size=(default_size or 11) - 0.5, color=DARK_BLUE)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "EDF1F5")
            run._r.get_or_add_rPr().append(shading)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_font(run, size=default_size, color=default_color)


def configure_section(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    left = p.add_run("GZWhisper")
    set_font(left, size=9, color=MUTED, bold=True)
    p.add_run("\t")
    right = p.add_run("Функциональный baseline для исследования конкурентов")
    set_font(right, size=9, color=MUTED)

    footer = section.footer
    p = footer.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    left = p.add_run("v1.5.2 + commit 05cdd4d · 03.09.2026")
    set_font(left, size=9, color=MUTED)
    p.add_run("\t")
    label = p.add_run("Стр. ")
    set_font(label, size=9, color=MUTED)
    add_page_field(p)


def add_cover(doc):
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(70)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(14)
    run = kicker.add_run("PRODUCT REFERENCE · DEEP RESEARCH INPUT")
    set_font(run, size=10, color=BLUE, bold=True)

    title = doc.add_paragraph(style="Cover Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("GZWhisper")

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Полное описание текущего функционала")

    purpose = doc.add_paragraph()
    purpose.alignment = WD_ALIGN_PARAGRAPH.CENTER
    purpose.paragraph_format.space_after = Pt(34)
    run = purpose.add_run("Контекст для глубокого исследования конкурентов\nи поиска функций для продуктового развития")
    set_font(run, size=12, color=DARK_BLUE, bold=True)

    callout = doc.add_paragraph()
    callout.paragraph_format.left_indent = Inches(0.55)
    callout.paragraph_format.right_indent = Inches(0.55)
    callout.paragraph_format.space_before = Pt(0)
    callout.paragraph_format.space_after = Pt(30)
    callout.paragraph_format.line_spacing = 1.2
    callout.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_cell_or_paragraph_shading(callout, LIGHT_FILL)
    set_left_border(callout)
    run = callout.add_run(
        "Local-first desktop-транскрибация · запись system audio и микрофона · "
        "управляемая очередь · локальная история · TXT/JSON"
    )
    set_font(run, size=11, color=INK, bold=True)

    metadata = doc.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata.paragraph_format.space_after = Pt(5)
    run = metadata.add_run("Снимок продукта: 3 сентября 2026 года")
    set_font(run, size=10.5, color=MUTED)
    metadata = doc.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata.paragraph_format.space_after = Pt(5)
    run = metadata.add_run("Версия: v1.5.2, build 270826")
    set_font(run, size=10.5, color=MUTED)
    metadata = doc.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata.paragraph_format.space_after = Pt(0)
    run = metadata.add_run("Учтён HEAD: 05cdd4d — Name transcript exports after history items")
    set_font(run, size=10.5, color=MUTED)

    doc.add_page_break()


def add_toc(doc, headings):
    heading = doc.add_paragraph("Оглавление", style="Heading 1")
    heading.paragraph_format.space_before = Pt(0)
    for title in headings:
        p = doc.add_paragraph(style="TOC Entry")
        add_inline_markdown(p, title, default_size=10.5, default_color=DARK_BLUE)
    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(10)
    add_inline_markdown(
        note,
        "Раздел 16 содержит готовую постановку задачи, которую можно передать ChatGPT Pro вместе с документом.",
        default_size=10.5,
        default_color=MUTED,
    )
    doc.add_page_break()


def build():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    headings = [line[3:].strip() for line in lines if line.startswith("## ")]

    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0])
    add_cover(doc)
    add_toc(doc, headings)

    started = False
    current_list_kind = None
    current_list_num = None
    for raw in lines:
        line = raw.rstrip()
        if line.startswith("## "):
            current_list_kind = None
            current_list_num = None
            started = True
            p = doc.add_paragraph(style="Heading 1")
            add_inline_markdown(p, line[3:].strip(), default_size=16, default_color=BLUE)
            continue
        if not started:
            continue
        if line.startswith("### "):
            current_list_kind = None
            current_list_num = None
            p = doc.add_paragraph(style="Heading 2")
            add_inline_markdown(p, line[4:].strip(), default_size=13, default_color=BLUE)
            continue
        if line.startswith("#### "):
            current_list_kind = None
            current_list_num = None
            p = doc.add_paragraph(style="Heading 3")
            add_inline_markdown(p, line[5:].strip(), default_size=12, default_color=DARK_BLUE)
            continue
        if not line or line == "---":
            current_list_kind = None
            current_list_num = None
            continue
        if line.startswith(">"):
            current_list_kind = None
            current_list_num = None
            content = line[1:].lstrip()
            if not content:
                continue
            p = doc.add_paragraph(style="Research Prompt")
            set_cell_or_paragraph_shading(p, LIGHT_FILL)
            set_left_border(p)
            add_inline_markdown(p, content, default_size=10.5, default_color=RGBColor(0x18, 0x1D, 0x24))
            continue
        bullet = re.match(r"^-\s+(.*)$", line)
        if bullet:
            if current_list_kind != "bullet":
                current_list_kind = "bullet"
                current_list_num = create_numbering(doc, "bullet")
            p = doc.add_paragraph()
            apply_numbering(p, current_list_num)
            add_inline_markdown(p, bullet.group(1))
            continue
        numbered = re.match(r"^\d+\.\s+(.*)$", line)
        if numbered:
            if current_list_kind != "decimal":
                current_list_kind = "decimal"
                current_list_num = create_numbering(doc, "decimal")
            p = doc.add_paragraph()
            apply_numbering(p, current_list_num)
            add_inline_markdown(p, numbered.group(1))
            continue
        current_list_kind = None
        current_list_num = None
        p = doc.add_paragraph()
        add_inline_markdown(p, line)

    core = doc.core_properties
    core.title = "GZWhisper: полное описание текущего функционала"
    core.subject = "Baseline для исследования конкурентов"
    core.author = "GZWhisper project"
    core.keywords = "GZWhisper, transcription, competitor research, local-first"

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    try:
        build()
    except Exception as exc:
        print(f"Failed to build document: {exc}", file=sys.stderr)
        raise
